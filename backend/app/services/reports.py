"""Accreditation-ready report generation (PDF via ReportLab, Excel via openpyxl)."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

INSTITUTION = "Student Assessment & Analytics Platform"
HEADER_FILL = PatternFill("solid", fgColor="157A50")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D8DEE5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fmt(v, nd=2):
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:.{nd}f}"
    return str(v)


# ------------------------------------------------------------------ PDF helpers

def _styles():
    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle("Cell", parent=ss["BodyText"], fontSize=8, leading=10))
    ss.add(ParagraphStyle("Small", parent=ss["BodyText"], fontSize=8, textColor=colors.HexColor("#555555")))
    return ss


def _table(data, col_widths=None, header_rows=1):
    t = Table(data, colWidths=col_widths, repeatRows=header_rows)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, header_rows - 1), colors.HexColor("#157A50")),
        ("TEXTCOLOR", (0, 0), (-1, header_rows - 1), colors.white),
        ("FONTNAME", (0, 0), (-1, header_rows - 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D8DEE5")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, header_rows), (-1, -1), [colors.white, colors.HexColor("#F1FAF5")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return t


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#777777"))
    canvas.drawString(15 * mm, 10 * mm, f"{INSTITUTION} · generated {datetime.now():%d %b %Y %H:%M}")
    canvas.drawRightString(doc.pagesize[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
    canvas.restoreState()


def attainment_pdf(att: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=14 * mm, bottomMargin=16 * mm,
                            title=f"CO-PO Attainment {att['course']['code']}")
    ss = _styles()
    c, cfg = att["course"], att["config"]
    el = [
        Paragraph("Course Outcome / Program Outcome Attainment Report", ss["Title"]),
        Paragraph(f"<b>{c['code']} — {c['name']}</b> · {c['program'] or ''} · Semester {c['semester']} · "
                  f"Section {c['section']} · Academic year {c['academic_year']} · Faculty: {c['faculty'] or '—'}",
                  ss["BodyText"]),
        Paragraph(f"Students evaluated: {att['students']} · CO target: {cfg['co_target_pct']:.0f}% marks · "
                  f"Levels: 3 ≥ {cfg['level3_pct']:.0f}%, 2 ≥ {cfg['level2_pct']:.0f}%, "
                  f"1 ≥ {cfg['level1_pct']:.0f}% of students · Weights CP {cfg['cp_weight']:.0f} / "
                  f"NCP {cfg['ncp_weight']:.0f} · Target attainment level {cfg['target_level']:.1f}",
                  ss["Small"]),
        Spacer(1, 6),
        Paragraph("1. Course Outcome attainment (direct)", ss["Heading3"]),
    ]
    rows = [["CO", "Statement", "CP assessments", "CP: % ≥ target", "CP level",
             "NCP: % ≥ target", "NCP level", "Attainment", "Status"]]
    for co in att["cos"]:
        rows.append([
            co["code"], Paragraph(co["description"], ss["Cell"]),
            Paragraph(", ".join(co["assessments"]) or "—", ss["Cell"]),
            _fmt(co["cp"]["pct_students"], 1), _fmt(co["cp"]["level"]),
            _fmt(co["ncp"]["pct_students"], 1), _fmt(co["ncp"]["level"]),
            _fmt(co["attainment"]),
            "—" if co["attained"] is None else ("Attained" if co["attained"] else "Not attained"),
        ])
    el.append(_table(rows, [14 * mm, 70 * mm, 50 * mm, 22 * mm, 16 * mm, 22 * mm, 16 * mm, 20 * mm, 24 * mm]))

    if att["pos"]:
        el += [Spacer(1, 8), Paragraph("2. CO–PO mapping matrix (1 = low, 2 = medium, 3 = high)", ss["Heading3"])]
        pos = att["pos"]
        mrows = [["CO"] + [p["code"] for p in pos]]
        for m in att["matrix"]:
            mrows.append([m["co"]] + [_fmt(m["values"].get(p["id"]) or m["values"].get(str(p["id"]))) for p in pos])
        mrows.append(["PO attainment"] + [_fmt(p["attainment"]) for p in pos])
        el.append(_table(mrows))
        el += [Spacer(1, 4), Paragraph(
            "PO attainment = Σ(CO attainment × mapping strength) / Σ(mapping strength).", ss["Small"])]
    doc.build(el, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


def student_report_pdf(student: dict, courses: list[dict]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=14 * mm, bottomMargin=16 * mm, title=f"Performance report {student['name']}")
    ss = _styles()
    el = [Paragraph("Student Performance Report", ss["Title"]),
          Paragraph(f"<b>{student['name']}</b> · Roll no. {student.get('roll_no') or '—'} · "
                    f"{student.get('program') or ''}", ss["BodyText"]), Spacer(1, 8)]
    for i, c in enumerate(courses):
        el.append(Paragraph(f"{c['code']} — {c['name']} ({c['academic_year']})", ss["Heading3"]))
        rows = [["Assessment", "Component", "Marks", "Max", "%", "Weight"]]
        for a in c["assessments"]:
            rows.append([Paragraph(a["title"], ss["Cell"]), a["component"], _fmt(a["marks"]),
                         _fmt(a["max_marks"], 0), _fmt(a["pct"], 1), _fmt(a["weightage"], 0)])
        el.append(_table(rows, [60 * mm, 28 * mm, 22 * mm, 18 * mm, 18 * mm, 20 * mm]))
        el.append(Spacer(1, 4))
        el.append(Paragraph(
            f"CP score: <b>{_fmt(c['cp_pct'])}%</b> · NCP: <b>{_fmt(c['ncp_pct'])}%</b> · "
            f"Final: <b>{_fmt(c['final_pct'])}%</b> · Grade: <b>{c['grade'] or '—'}</b>", ss["BodyText"]))
        if c.get("risk"):
            el.append(Paragraph(f"Early-warning status: {c['risk']['risk_level'].upper()}", ss["Small"]))
        el.append(Spacer(1, 8))
        if i and i % 3 == 0:
            el.append(PageBreak())
    doc.build(el, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


# ------------------------------------------------------------------ Excel helpers

def _header(ws, row: int, values: list):
    for i, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, max_width=60):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_width, max(10, width + 2))


def _xlsx_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def attainment_xlsx(att: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "CO Attainment"
    c, cfg = att["course"], att["config"]
    ws.append([f"CO/PO Attainment — {c['code']} {c['name']}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Academic year {c['academic_year']} · Semester {c['semester']} · Section {c['section']} · "
               f"Faculty {c['faculty'] or '—'} · Students {att['students']}"])
    ws.append([f"CO target {cfg['co_target_pct']}% · L3 ≥ {cfg['level3_pct']}% · L2 ≥ {cfg['level2_pct']}% · "
               f"L1 ≥ {cfg['level1_pct']}% · CP/NCP weights {cfg['cp_weight']}/{cfg['ncp_weight']}"])
    _header(ws, 5, ["CO", "Statement", "CP assessments", "CP evaluated", "CP attained", "CP % students",
                    "CP level", "NCP % students", "NCP level", "Attainment", "Status"])
    for co in att["cos"]:
        ws.append([co["code"], co["description"], ", ".join(co["assessments"]), co["cp"]["evaluated"],
                   co["cp"]["attained"], co["cp"]["pct_students"], co["cp"]["level"],
                   co["ncp"]["pct_students"], co["ncp"]["level"], co["attainment"],
                   None if co["attained"] is None else ("Attained" if co["attained"] else "Not attained")])
    _autosize(ws)

    if att["pos"]:
        ws2 = wb.create_sheet("CO-PO Matrix")
        pos = att["pos"]
        _header(ws2, 1, ["CO"] + [p["code"] for p in pos])
        for m in att["matrix"]:
            ws2.append([m["co"]] + [m["values"].get(p["id"]) for p in pos])
        ws2.append(["PO attainment"] + [p["attainment"] for p in pos])
        for cell in ws2[ws2.max_row]:
            cell.font = Font(bold=True)
        _autosize(ws2)
    return _xlsx_bytes(wb)


def gradebook_xlsx(course, gb: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidated CP-NCP"
    ws.append([f"{course.code} {course.name} — consolidated CP–NCP result ({course.academic_year})"])
    ws["A1"].font = Font(bold=True, size=13)
    heads = ["Roll no", "Name", "Email"]
    heads += [f"{a['title']} (/{a['max_marks']:g}, {a['weightage']:g}%)" for a in gb["assessments"]]
    heads += ["Quiz %", "Assignment %", "Practical %", "Attendance %", "CP %",
              f"NCP (/{course.ncp_max_marks:g})", "NCP %", "Final %", "Grade", "Result"]
    _header(ws, 3, heads)
    for s in gb["students"]:
        row = [s["roll_no"], s["name"], s["email"]]
        for a in gb["assessments"]:
            cell = s["assessments"].get(a["id"]) or {}
            row.append("AB" if cell.get("absent") else cell.get("marks"))
        comp = s["components"]
        row += [comp.get("quiz"), comp.get("assignment"), comp.get("practical"), comp.get("attendance"),
                s["cp_pct"], "AB" if s["ncp_absent"] else s["ncp_marks"], s["ncp_pct"], s["final_pct"],
                s["grade"], None if s["passed"] is None else ("PASS" if s["passed"] else "FAIL")]
        ws.append(row)
    ws.freeze_panes = "D4"
    _autosize(ws, 28)
    return _xlsx_bytes(wb)


def program_xlsx(program, data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "PO Attainment"
    ws.append([f"Program Outcome attainment — {program.code} {program.name}"
               + (f" ({data['academic_year']})" if data.get("academic_year") else "")])
    ws["A1"].font = Font(bold=True, size=13)
    pos = data["po_list"]
    _header(ws, 3, ["Course", "Name", "Sem", "Year"] + [p["code"] for p in pos])
    for c in data["courses"]:
        ws.append([c["code"], c["name"], c["semester"], c["academic_year"]]
                  + [c["values"].get(p["id"]) for p in pos])
    ws.append(["Program average", "", "", ""] + [p["attainment"] for p in data["pos"]])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    _autosize(ws, 40)
    return _xlsx_bytes(wb)
